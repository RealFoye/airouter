"""AirRouter 用户工作台：先分析，再执行。"""

from __future__ import annotations

import base64
import csv
import html
import json
import re
import uuid
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
from time import perf_counter
from typing import Any

import streamlit as st
from openai import APIError, APITimeoutError, AuthenticationError, OpenAI, RateLimitError

from airouter.adapters.base import (
    AdapterAuthError,
    AdapterError,
    AdapterRateLimitError,
    AdapterTimeoutError,
)
from airouter.adapters.openai_compat import OpenAICompatAdapter
from airouter.config import load_all
from airouter.engine.complexity.types import ComplexityContext
from airouter.engine.router.router import Router
from airouter.models.enums import (
    ComplexityLevel,
    RiskLevel,
    StrategyType,
    TaskCategory,
    VerifiabilityLevel,
)
from airouter.models.model_registry import ModelRegistry
from airouter.models.schemas import RoutingDecision
from airouter.web.data import model_catalog
from airouter.web.ui import configure_page, require_login


@dataclass(frozen=True)
class ModelBenchmark:
    pass_rate: float
    avg_score: float
    cost_cny: float
    position: str


MODEL_BENCHMARKS: dict[str, ModelBenchmark] = {
    "qwen3.7-plus": ModelBenchmark(72.2, 0.826, 1.48, "质量最稳"),
    "deepseek-v4-pro": ModelBenchmark(70.0, 0.800, 0.59, "性价比最佳"),
    "glm-5.1": ModelBenchmark(65.6, 0.800, 4.19, "强模型备选"),
    "qwen3.7-max": ModelBenchmark(64.4, 0.799, 8.16, "高成本强模型"),
    "deepseek-v4-flash": ModelBenchmark(64.4, 0.786, 0.20, "成本最低"),
}

MODEL_DESCRIPTIONS = {
    "deepseek-v4-flash": "便宜、响应快，适合简单和容易检查的任务。",
    "deepseek-v4-pro": "质量接近高档模型，成本更低，适合作为默认稳妥选择。",
    "qwen3.7-plus": "综合质量最高，适合需要更稳结果的任务。",
    "qwen3.7-max": "成本较高，适合特别需要强模型兜底的任务。",
    "glm-5.1": "强模型备选，适合需要交叉验证或替代方案时使用。",
}

SYSTEM_PROMPT = (
    "你是企业工作助手。请严格按用户任务执行。"
    "如果信息不足，先说明需要补充什么；"
    "不要声称读取、运行或修改了你没有实际操作过的文件、测试或系统。"
)
CONTEXT_COMPRESSION_CHAR_LIMIT = 16000
CONTEXT_KEEP_RECENT_MESSAGES = 8
CONTEXT_SUMMARY_MODEL = "qwen3.7-max"
CONVERSATION_TITLE_MODEL = "qwen3.7-max"
RECOMMENDATION_REVIEW_MODEL = "MiniMax-M3"
RECOMMENDATION_REVIEW_FALLBACK_MODEL = "qwen3.7-max"
IMAGE_MODEL = "qwen3.7-plus"
MAX_ATTACHMENTS = 4
MAX_ATTACHMENT_BYTES = 4 * 1024 * 1024
MAX_TEXT_ATTACHMENT_CHARS = 12000
CONVERSATION_STORE_PATH = Path("runs/web/conversations.json")
OPENSQUILLA_BUNDLE_PATH = Path(
    "opensquilla/src/opensquilla/squilla_router/models/v4.2_phase3_inference"
)
TEXT_ATTACHMENT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".tsv",
    ".json",
    ".jsonl",
    ".log",
    ".py",
    ".sql",
    ".yaml",
    ".yml",
    ".xml",
    ".html",
    ".css",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
}
SPREADSHEET_ATTACHMENT_EXTENSIONS = {".xlsx", ".xlsm"}
UPLOAD_FILE_TYPES = [
    "png",
    "jpg",
    "jpeg",
    "webp",
    "txt",
    "md",
    "csv",
    "tsv",
    "json",
    "jsonl",
    "log",
    "py",
    "sql",
    "yaml",
    "yml",
    "xml",
    "html",
    "css",
    "js",
    "jsx",
    "ts",
    "tsx",
    "xlsx",
    "xlsm",
    "pdf",
]
PROCESSING_MODES = ["自动", "只读取文件", "生成报告"]
REPORT_REQUEST_TERMS = (
    "总结",
    "报告",
    "日报",
    "周报",
    "月报",
    "会议纪要",
    "纪要",
    "整理成文档",
    "整理文档",
    "项目进展",
    "复盘",
    "汇报",
)


def _blank_conversation(title: str = "新对话") -> dict[str, Any]:
    return {
        "title": title,
        "messages": [
            {
                "role": "assistant",
                "content": "告诉我你想处理什么。我会先分析该用哪个模型和处理方式，再由你决定是否开始执行。",
            }
        ],
        "analysis": None,
        "selected_model": None,
        "selected_treatment": None,
        "execution_started": False,
        "manual_open": False,
        "execution_records": [],
        "execution_messages": [],
        "context_summary": "",
        "context_summary_updated_at": "",
        "judge_review": None,
        "attachments": [],
    }


def _serialize_analysis(analysis: Any) -> dict[str, Any] | None:
    if not analysis:
        return None
    decision = analysis.get("decision") if isinstance(analysis, dict) else None
    if isinstance(decision, RoutingDecision):
        decision_payload: dict[str, Any] | None = decision.model_dump(mode="json")
    elif isinstance(decision, dict):
        decision_payload = decision
    else:
        decision_payload = None
    return {
        "task": analysis.get("task", "") if isinstance(analysis, dict) else "",
        "decision": decision_payload,
    }


def _deserialize_analysis(analysis: Any) -> dict[str, Any] | None:
    if not isinstance(analysis, dict) or not analysis.get("decision"):
        return None
    try:
        decision = RoutingDecision.model_validate(analysis["decision"])
    except Exception:
        return None
    return {
        "task": str(analysis.get("task", "")),
        "decision": decision,
    }


def _serialize_conversation(conversation: dict[str, Any]) -> dict[str, Any]:
    data = dict(conversation)
    data["analysis"] = _serialize_analysis(data.get("analysis"))
    return data


def _deserialize_conversation(conversation: dict[str, Any]) -> dict[str, Any]:
    data = _blank_conversation(str(conversation.get("title") or "新对话"))
    data.update(conversation)
    data["analysis"] = _deserialize_analysis(conversation.get("analysis"))
    return data


def _load_conversation_store() -> dict[str, Any]:
    if not CONVERSATION_STORE_PATH.exists():
        return {}
    try:
        raw = json.loads(CONVERSATION_STORE_PATH.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}
    return raw if isinstance(raw, dict) else {}


def _save_conversation_store(store: dict[str, Any]) -> None:
    CONVERSATION_STORE_PATH.parent.mkdir(parents=True, exist_ok=True)
    CONVERSATION_STORE_PATH.write_text(
        json.dumps(store, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _current_user_key() -> str:
    return str(st.session_state.get("api_key_fingerprint") or "anonymous")


def _load_user_conversations() -> dict[str, Any]:
    store = _load_conversation_store()
    bucket = store.get(_current_user_key(), {})
    conversations = bucket.get("conversations", {}) if isinstance(bucket, dict) else {}
    if not isinstance(conversations, dict):
        return {}
    return {
        conv_id: _deserialize_conversation(conversation)
        for conv_id, conversation in conversations.items()
        if isinstance(conversation, dict)
    }


def _persist_user_conversations() -> None:
    user_key = _current_user_key()
    store = _load_conversation_store()
    conversations = st.session_state.get("ar_conversations", {})
    store[user_key] = {
        "updated_at": datetime.now().isoformat(timespec="seconds"),
        "current_conversation_id": st.session_state.get("ar_current_conversation_id"),
        "conversations": {
            conv_id: _serialize_conversation(conversation)
            for conv_id, conversation in conversations.items()
        },
    }
    _save_conversation_store(store)


def _init_state() -> None:
    user_key = _current_user_key()
    if st.session_state.get("ar_loaded_user_key") != user_key:
        for key in list(st.session_state.keys()):
            if key.startswith("ar_"):
                st.session_state.pop(key, None)
        saved = _load_user_conversations()
        if saved:
            store = _load_conversation_store()
            bucket = store.get(user_key, {})
            current_id = (
                bucket.get("current_conversation_id")
                if isinstance(bucket, dict)
                else None
            )
            if current_id not in saved:
                current_id = next(iter(saved))
            st.session_state["ar_current_conversation_id"] = current_id
            st.session_state["ar_conversations"] = saved
        st.session_state["ar_loaded_user_key"] = user_key

    if "ar_conversations" not in st.session_state:
        conv_id = uuid.uuid4().hex[:10]
        st.session_state["ar_current_conversation_id"] = conv_id
        st.session_state["ar_conversations"] = {conv_id: _blank_conversation()}
    current_id = st.session_state["ar_current_conversation_id"]
    current = st.session_state["ar_conversations"][current_id]
    st.session_state.setdefault("ar_messages", current["messages"])
    st.session_state.setdefault("ar_analysis", current["analysis"])
    st.session_state.setdefault("ar_selected_model", current["selected_model"])
    st.session_state.setdefault("ar_selected_treatment", current["selected_treatment"])
    st.session_state.setdefault("ar_execution_started", current["execution_started"])
    st.session_state.setdefault("ar_manual_open", current["manual_open"])
    st.session_state.setdefault("ar_execution_records", current["execution_records"])
    st.session_state.setdefault("ar_execution_messages", current["execution_messages"])
    st.session_state.setdefault("ar_context_summary", current["context_summary"])
    st.session_state.setdefault(
        "ar_context_summary_updated_at",
        current["context_summary_updated_at"],
    )
    st.session_state.setdefault("ar_judge_review", current.get("judge_review"))
    st.session_state.setdefault("ar_attachments", current.get("attachments", []))


def _save_current_conversation() -> None:
    current_id = st.session_state.get("ar_current_conversation_id")
    if not current_id:
        return
    conversations = st.session_state.get("ar_conversations", {})
    if current_id not in conversations:
        return
    conversations[current_id].update(
        {
            "messages": list(st.session_state.get("ar_messages", [])),
            "analysis": st.session_state.get("ar_analysis"),
            "selected_model": st.session_state.get("ar_selected_model"),
            "selected_treatment": st.session_state.get("ar_selected_treatment"),
            "execution_started": st.session_state.get("ar_execution_started", False),
            "manual_open": st.session_state.get("ar_manual_open", False),
            "execution_records": list(st.session_state.get("ar_execution_records", [])),
            "execution_messages": list(st.session_state.get("ar_execution_messages", [])),
            "context_summary": st.session_state.get("ar_context_summary", ""),
            "context_summary_updated_at": st.session_state.get(
                "ar_context_summary_updated_at", ""
            ),
            "judge_review": st.session_state.get("ar_judge_review"),
            "attachments": list(st.session_state.get("ar_attachments", [])),
        }
    )
    _persist_user_conversations()


def _load_conversation(conversation_id: str) -> None:
    _save_current_conversation()
    current = st.session_state["ar_conversations"][conversation_id]
    st.session_state["ar_current_conversation_id"] = conversation_id
    st.session_state["ar_messages"] = list(current["messages"])
    st.session_state["ar_analysis"] = current["analysis"]
    st.session_state["ar_selected_model"] = current["selected_model"]
    st.session_state["ar_selected_treatment"] = current["selected_treatment"]
    st.session_state["ar_execution_started"] = current["execution_started"]
    st.session_state["ar_manual_open"] = current["manual_open"]
    st.session_state["ar_execution_records"] = list(current["execution_records"])
    st.session_state["ar_execution_messages"] = list(current["execution_messages"])
    st.session_state["ar_context_summary"] = current["context_summary"]
    st.session_state["ar_context_summary_updated_at"] = current["context_summary_updated_at"]
    st.session_state["ar_judge_review"] = current.get("judge_review")
    st.session_state["ar_attachments"] = list(current.get("attachments", []))


def _switch_to_conversation(conversation_id: str) -> None:
    current = st.session_state["ar_conversations"][conversation_id]
    st.session_state["ar_current_conversation_id"] = conversation_id
    st.session_state["ar_messages"] = list(current["messages"])
    st.session_state["ar_analysis"] = current["analysis"]
    st.session_state["ar_selected_model"] = current["selected_model"]
    st.session_state["ar_selected_treatment"] = current["selected_treatment"]
    st.session_state["ar_execution_started"] = current["execution_started"]
    st.session_state["ar_manual_open"] = current["manual_open"]
    st.session_state["ar_execution_records"] = list(current["execution_records"])
    st.session_state["ar_execution_messages"] = list(current["execution_messages"])
    st.session_state["ar_context_summary"] = current["context_summary"]
    st.session_state["ar_context_summary_updated_at"] = current["context_summary_updated_at"]
    st.session_state["ar_judge_review"] = current.get("judge_review")
    st.session_state["ar_attachments"] = list(current.get("attachments", []))


def _new_conversation() -> None:
    _save_current_conversation()
    conv_id = uuid.uuid4().hex[:10]
    st.session_state["ar_conversations"][conv_id] = _blank_conversation()
    _load_conversation(conv_id)


def _fallback_conversation_title(task: str) -> str:
    title = re.sub(r"\s+", " ", task.strip())
    title = re.sub(r"^#+\s*", "", title)
    if len(title) > 18:
        title = title[:18] + "..."
    return title or "新对话"


def _clean_conversation_title(raw_title: str) -> str:
    title = _strip_thinking_text(raw_title)
    title = re.sub(r"```(?:text|markdown)?", "", title, flags=re.IGNORECASE)
    title = title.replace("```", "")
    title = title.splitlines()[0] if title.splitlines() else title
    title = re.sub(r"^\s*(标题|对话标题)\s*[:：]\s*", "", title)
    title = title.strip(" \t\r\n\"'“”‘’`。、，,.;；:")
    title = re.sub(r"\s+", "", title)
    if len(title) > 16:
        title = title[:16] + "..."
    return title


def _clean_manual_conversation_title(raw_title: str) -> str:
    title = _strip_thinking_text(raw_title)
    title = title.splitlines()[0] if title.splitlines() else title
    title = title.strip(" \t\r\n\"'“”‘’`。、，,.;；:")
    title = re.sub(r"\s+", " ", title)
    if len(title) > 24:
        title = title[:24] + "..."
    return title or "新对话"


def _generate_conversation_title(task: str, api_key: str) -> str:
    fallback = _fallback_conversation_title(task)
    if not api_key:
        return fallback
    try:
        model = _model_entry(CONVERSATION_TITLE_MODEL)
        adapter = OpenAICompatAdapter(api_base=model.api_base, api_key=api_key)
        response = adapter.chat(
            CONVERSATION_TITLE_MODEL,
            [
                {
                    "role": "system",
                    "content": (
                        "你负责给企业 AI 助手对话生成中文短标题。"
                        "只输出标题本身，不要解释，不要 Markdown，不要引号。"
                    ),
                },
                {
                    "role": "user",
                    "content": (
                        "请根据下面的用户问题生成一个 6 到 12 个中文字的对话标题。"
                        "标题要具体，不要写“新对话”“问题分析”这种泛泛词。\n\n"
                        f"{task.strip()[:1800]}"
                    ),
                },
            ],
            temperature=0.0,
            max_tokens=40,
        )
    except (AdapterAuthError, AdapterRateLimitError, AdapterTimeoutError, AdapterError):
        return fallback
    return _clean_conversation_title(response.content) or fallback


def _update_conversation_title(task: str, api_key: str | None = None) -> None:
    current_id = st.session_state.get("ar_current_conversation_id")
    if not current_id:
        return
    title = _generate_conversation_title(task, api_key or "")
    st.session_state["ar_conversations"][current_id]["title"] = title or "新对话"


def _rename_conversation(conversation_id: str, title: str) -> None:
    conversations = st.session_state.get("ar_conversations", {})
    if conversation_id not in conversations:
        return
    conversations[conversation_id]["title"] = _clean_manual_conversation_title(title)
    _persist_user_conversations()


def _delete_conversation(conversation_id: str) -> None:
    conversations = st.session_state.get("ar_conversations", {})
    if conversation_id not in conversations:
        return
    was_current = conversation_id == st.session_state.get("ar_current_conversation_id")
    conversations.pop(conversation_id, None)
    if not conversations:
        new_id = uuid.uuid4().hex[:10]
        conversations[new_id] = _blank_conversation()
    if was_current:
        next_id = next(reversed(conversations))
        _switch_to_conversation(next_id)
    _persist_user_conversations()


def _render_conversation_sidebar() -> None:
    with st.sidebar:
        st.divider()
        st.markdown("### 对话")
        if st.button("新对话", use_container_width=True):
            _new_conversation()
            st.rerun()

        conversations = st.session_state.get("ar_conversations", {})
        current_id = st.session_state.get("ar_current_conversation_id")
        for conv_id, conversation in reversed(list(conversations.items())):
            title = str(conversation.get("title") or "新对话")
            prefix = "当前 · " if conv_id == current_id else ""
            title_col, menu_col = st.columns([0.86, 0.14], gap="small")
            with title_col:
                if st.button(
                    f"{prefix}{title}",
                    key=f"conversation_{conv_id}",
                    use_container_width=True,
                ):
                    if conv_id != current_id:
                        _load_conversation(conv_id)
                        st.rerun()
            with menu_col:
                with st.popover(">", use_container_width=True):
                    new_title = st.text_input(
                        "重命名",
                        value=title,
                        key=f"rename_input_{conv_id}",
                    )
                    if st.button(
                        "保存",
                        key=f"rename_save_{conv_id}",
                        use_container_width=True,
                    ):
                        _rename_conversation(conv_id, new_title)
                        st.rerun()
                    if st.button(
                        "删除",
                        key=f"delete_conversation_{conv_id}",
                        use_container_width=True,
                    ):
                        _delete_conversation(conv_id)
                        st.rerun()


def _html_text(text: str) -> str:
    return html.escape(text).replace("\n", "<br>")


def _safe_attachment_name(name: str) -> str:
    cleaned = Path(name).name.strip()
    return cleaned or "attachment"


def _attachment_label(kind: str) -> str:
    return {
        "image": "图片",
        "text": "文本",
        "spreadsheet": "表格",
        "pdf": "PDF",
        "unsupported": "附件",
    }.get(kind, "附件")


def _processing_mode_instruction(mode: str) -> str:
    if mode == "只读取文件":
        return (
            "处理方式：只读取文件。请优先基于用户上传或粘贴的材料回答，"
            "不要主动写成长篇报告，除非用户明确要求。"
        )
    if mode == "生成报告":
        return (
            "处理方式：生成报告。请将结果整理为中文结构化报告，"
            "包含标题、核心结论、关键信息、事实依据、问题与建议。"
        )
    return ""


def _should_use_report_tool(mode: str, *texts: str) -> bool:
    if mode == "生成报告":
        return True
    if mode == "只读取文件":
        return False
    combined = "\n".join(texts).lower()
    return any(term.lower() in combined for term in REPORT_REQUEST_TERMS)


def _report_tool_instruction() -> str:
    return (
        "工具：报告整理。请将结果整理为中文结构化报告，"
        "包含标题、核心结论、关键信息、事实依据、问题与建议；"
        "如果用户要求日报/周报，则使用“完成内容、收获与感悟、下一步计划”的结构。"
    )


def _tool_status_items(
    *,
    attachments: list[dict[str, Any]] | None = None,
    uploaded_count: int = 0,
    report_tool: bool = False,
) -> list[dict[str, str]]:
    items: list[dict[str, str]] = []
    attachment_count = len(attachments or [])
    if attachment_count or uploaded_count:
        count = attachment_count or uploaded_count
        items.append(
            {
                "name": "文件解析",
                "detail": f"读取 {count} 个附件，提取可用正文或图片内容",
            }
        )
    if report_tool:
        items.append(
            {
                "name": "报告整理",
                "detail": "按中文报告结构组织输出",
            }
        )
    return items


def _tool_status_message(items: list[dict[str, str]], *, active: bool = False) -> str:
    if not items:
        return ""
    prefix = "正在使用工具" if active else "将使用工具"
    lines = [f"{prefix}："]
    lines.extend(f"- {item['name']}：{item['detail']}" for item in items)
    return "\n".join(lines)


def _render_tool_strip(items: list[dict[str, str]], *, active: bool = False) -> None:
    if not items:
        return
    label = "正在使用工具" if active else "将使用工具"
    chips = "".join(
        f"<span><strong>{html.escape(item['name'])}</strong>{html.escape(item['detail'])}</span>"
        for item in items
    )
    st.markdown(
        f'<div class="ar-tool-strip"><em>{label}</em>{chips}</div>',
        unsafe_allow_html=True,
    )


def _render_attachment_chips(attachments: list[dict[str, Any]]) -> None:
    if not attachments:
        return
    chips = []
    for attachment in attachments:
        label = _attachment_label(str(attachment.get("kind", "attachment")))
        name = html.escape(str(attachment.get("name") or "attachment"))
        chips.append(f"<span>{label} · {name}</span>")
    st.markdown(
        '<div class="ar-attachment-chips">' + "".join(chips) + "</div>",
        unsafe_allow_html=True,
    )


def _render_material_popover(
    key_prefix: str,
    *,
    retained_attachments: list[dict[str, Any]] | None = None,
) -> tuple[list[Any], str]:
    with st.popover("＋ 添加文件", use_container_width=True):
        st.markdown("拖拽文件到这里，或点击上传。")
        uploaded_files = st.file_uploader(
            "上传图片、表格、PDF 或文本附件",
            type=UPLOAD_FILE_TYPES,
            accept_multiple_files=True,
            key=f"{key_prefix}_files",
            help="图片会使用 qwen3.7-plus 处理；CSV/Excel/PDF/文本会尽量提取正文作为上下文。",
        )
        pasted_context = st.text_area(
            "粘贴补充材料",
            placeholder="也可以直接粘贴日志、代码、表格说明、制度片段或业务背景。",
            height=96,
            key=f"{key_prefix}_context",
        )
        st.caption("单个附件最多 4MB，最多读取 4 个。上传后会显示在输入框下方。")
        if retained_attachments:
            st.markdown("已保留资料")
            _render_attachment_chips(retained_attachments)
            st.caption("这些资料会保留在当前对话中，后续可以继续引用。")
    return list(uploaded_files or []), pasted_context


def _truncate_attachment_text(text: str) -> str:
    normalized = text.replace("\x00", "").strip()
    if len(normalized) <= MAX_TEXT_ATTACHMENT_CHARS:
        return normalized
    return normalized[:MAX_TEXT_ATTACHMENT_CHARS] + "\n\n[附件内容较长，已截断]"


def _decode_text_bytes(raw: bytes) -> tuple[str, str | None]:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding), None
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "文件编码无法精确识别，已按替换字符方式读取。"


def _extract_delimited_table(raw: bytes, *, delimiter: str | None = None) -> tuple[str, str | None]:
    decoded, warning = _decode_text_bytes(raw)
    sample = decoded[:4096]
    if delimiter is None:
        try:
            dialect = csv.Sniffer().sniff(sample)
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = "\t" if "\t" in sample else ","
    reader = csv.reader(StringIO(decoded), delimiter=delimiter)
    rows: list[str] = []
    row_count = 0
    for row in reader:
        row_count += 1
        if row_count > 40:
            break
        rows.append(" | ".join(str(cell).strip() for cell in row[:16]))
    suffix = "" if row_count <= 40 else "\n[表格行数较多，仅提取前 40 行]"
    return _truncate_attachment_text("\n".join(rows) + suffix), warning


def _extract_spreadsheet(raw: bytes) -> tuple[str, str | None]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "", "当前环境缺少 openpyxl，暂时无法读取 Excel 内容。"

    try:
        workbook = load_workbook(BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:
        return "", f"Excel 解析失败：{exc}"

    sections: list[str] = []
    for sheet_index, sheet in enumerate(workbook.worksheets[:3], start=1):
        lines = [f"[Sheet {sheet_index}: {sheet.title}]"]
        for row_index, row in enumerate(
            sheet.iter_rows(max_row=40, max_col=16, values_only=True),
            start=1,
        ):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                lines.append(f"{row_index}: " + " | ".join(values))
        sections.append("\n".join(lines))
    warning = None
    if len(workbook.worksheets) > 3:
        warning = "Excel 工作表较多，仅提取前 3 个工作表。"
    return _truncate_attachment_text("\n\n".join(sections)), warning


def _extract_pdf(raw: bytes) -> tuple[str, str | None]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return "", "当前环境缺少 pypdf，暂时无法读取 PDF 正文。"

    try:
        reader = PdfReader(BytesIO(raw))
        pages = []
        for page_index, page in enumerate(reader.pages[:6], start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"[PDF 第 {page_index} 页]\n{text.strip()}")
    except Exception as exc:
        return "", f"PDF 解析失败：{exc}"

    warning = None
    if len(reader.pages) > 6:
        warning = "PDF 页数较多，仅提取前 6 页。"
    if not pages:
        warning = warning or "没有从 PDF 中提取到可读文本，可能是扫描件或图片型 PDF。"
    return _truncate_attachment_text("\n\n".join(pages)), warning


def _process_uploaded_files(uploaded_files: list[Any]) -> tuple[list[dict[str, Any]], list[str]]:
    attachments: list[dict[str, Any]] = []
    warnings: list[str] = []
    if len(uploaded_files) > MAX_ATTACHMENTS:
        warnings.append(f"最多读取 {MAX_ATTACHMENTS} 个附件，后续文件已忽略。")

    for uploaded in uploaded_files[:MAX_ATTACHMENTS]:
        name = _safe_attachment_name(str(getattr(uploaded, "name", "attachment")))
        raw = uploaded.getvalue()
        size = len(raw)
        mime_type = str(getattr(uploaded, "type", "") or "")
        suffix = Path(name).suffix.lower()
        if size > MAX_ATTACHMENT_BYTES:
            warnings.append(f"{name} 超过 4MB，已忽略。")
            continue

        base: dict[str, Any] = {
            "name": name,
            "mime_type": mime_type,
            "size": size,
        }
        text = ""
        warning: str | None = None

        if mime_type.startswith("image/") or suffix in {".png", ".jpg", ".jpeg", ".webp"}:
            attachments.append(
                {
                    **base,
                    "kind": "image",
                    "data_url": f"data:{mime_type or 'image/png'};base64,{base64.b64encode(raw).decode('ascii')}",
                    "text": "",
                }
            )
            continue

        if suffix in {".csv", ".tsv"}:
            text, warning = _extract_delimited_table(
                raw,
                delimiter="\t" if suffix == ".tsv" else None,
            )
            kind = "spreadsheet"
        elif suffix in SPREADSHEET_ATTACHMENT_EXTENSIONS:
            text, warning = _extract_spreadsheet(raw)
            kind = "spreadsheet"
        elif suffix == ".pdf" or mime_type == "application/pdf":
            text, warning = _extract_pdf(raw)
            kind = "pdf"
        elif suffix in TEXT_ATTACHMENT_EXTENSIONS or mime_type.startswith("text/"):
            decoded, warning = _decode_text_bytes(raw)
            text = _truncate_attachment_text(decoded)
            kind = "text"
        else:
            kind = "unsupported"
            warning = "暂不支持读取该附件正文，只会把文件名和类型作为上下文。"

        if warning:
            warnings.append(f"{name}: {warning}")
        attachments.append({**base, "kind": kind, "text": text})
    return attachments, warnings


def _has_image_attachments(attachments: list[dict[str, Any]] | None = None) -> bool:
    items = attachments if attachments is not None else st.session_state.get("ar_attachments", [])
    return any(item.get("kind") == "image" and item.get("data_url") for item in items)


def _image_attachments_only(attachments: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [item for item in attachments if item.get("kind") == "image" and item.get("data_url")]


def _attachment_summary_lines(attachments: list[dict[str, Any]]) -> list[str]:
    lines = []
    for attachment in attachments:
        kind = _attachment_label(str(attachment.get("kind", "attachment")))
        size = int(attachment.get("size") or 0)
        size_label = f"{size / 1024:.1f}KB" if size else "未知大小"
        lines.append(f"- {attachment.get('name', 'attachment')}（{kind}，{size_label}）")
    return lines


def _attachment_text_block(attachments: list[dict[str, Any]]) -> str:
    sections = []
    for attachment in attachments:
        name = str(attachment.get("name") or "attachment")
        kind = str(attachment.get("kind") or "attachment")
        label = _attachment_label(kind)
        text = str(attachment.get("text") or "").strip()
        if kind == "image":
            body = "这是一张图片，执行时会随请求发送给支持图片的模型。"
        elif text:
            body = text
        else:
            body = "未提取到正文内容，请不要假装已经读取文件内部细节。"
        sections.append(f"### 附件：{name}（{label}）\n{body}")
    return "\n\n".join(sections)


def _build_task_with_materials(
    task_text: str,
    extra_context: str,
    attachments: list[dict[str, Any]],
) -> str:
    return "\n\n".join(
        part
        for part in [
            task_text.strip(),
            extra_context.strip(),
            _attachment_text_block(attachments),
        ]
        if part
    )


def _build_user_display_text(task_text: str, attachments: list[dict[str, Any]]) -> str:
    lines = _attachment_summary_lines(attachments)
    if not lines:
        return task_text.strip()
    return f"{task_text.strip()}\n\n附件：\n" + "\n".join(lines)


def _build_multimodal_user_content(
    user_content: str,
    attachments: list[dict[str, Any]],
) -> str | list[dict[str, Any]]:
    material_block = _attachment_text_block(attachments)
    text = "\n\n".join(part for part in [user_content.strip(), material_block] if part)
    if not _has_image_attachments(attachments):
        return text
    parts: list[dict[str, Any]] = [{"type": "text", "text": text}]
    for attachment in attachments:
        if attachment.get("kind") == "image" and attachment.get("data_url"):
            parts.append(
                {
                    "type": "image_url",
                    "image_url": {"url": str(attachment["data_url"])},
                }
            )
    return parts


def _message_text_for_estimation(message: dict[str, Any]) -> str:
    content = message.get("content", "")
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        parts = []
        for item in content:
            if not isinstance(item, dict):
                continue
            if item.get("type") == "text":
                parts.append(str(item.get("text", "")))
            elif item.get("type") == "image_url":
                parts.append("[image_attachment]")
        return "\n".join(parts)
    return str(content)


def _friendly_strategy(decision: RoutingDecision) -> tuple[str, str]:
    if decision.requires_human_review:
        return (
            "需要确认后执行",
            "这类任务可能影响正式结果，建议你确认方案后再继续。",
        )
    if decision.recommended_strategy == StrategyType.PLAN_EXECUTE_REVIEW:
        return (
            "先计划、执行后复核",
            "系统会先拆解步骤，再执行，并在完成后做一次检查。",
        )
    if decision.recommended_strategy == StrategyType.PLAN_EXECUTE:
        return (
            "先计划再执行",
            "任务有多个步骤，先列清楚计划会更稳。",
        )
    if decision.recommended_strategy == StrategyType.GUIDED_EXECUTE:
        return (
            "稳妥处理",
            "任务需要一定推理和整理，使用更可靠的执行方式。",
        )
    if decision.task_features.complexity in (ComplexityLevel.SIMPLE, ComplexityLevel.MODERATE):
        return (
            "快速处理",
            "任务边界比较清楚，可以用较轻量的方式完成。",
        )
    return (
        "稳妥处理",
        "任务有一定复杂度，建议使用更可靠的模型完成。",
    )


def _primary_model(decision: RoutingDecision) -> str:
    return decision.models.executor or decision.models.planner or decision.models.reviewer or "人工确认"


def _completion_label(model_name: str, decision: RoutingDecision) -> str:
    stats = MODEL_BENCHMARKS.get(model_name)
    if stats:
        return f"约 {stats.pass_rate:.0f}%"
    quality = decision.estimated_quality
    if quality >= 0.85:
        return "较高"
    if quality >= 0.72:
        return "中等偏高"
    if quality >= 0.6:
        return "中等"
    return "需要谨慎"


def _cost_comparison(model_name: str, decision: RoutingDecision) -> str:
    payload = _cost_comparison_payload(model_name, decision)
    return str(payload["summary"])


def _cost_comparison_payload(model_name: str, decision: RoutingDecision) -> dict[str, str | float]:
    recommended_cost = _estimate_decision_cost_for_model(model_name, decision)
    max_model = "qwen3.7-max"
    max_cost = _estimate_decision_cost_for_model(max_model, decision)
    if model_name == max_model:
        return {
            "summary": "已是最高档",
            "detail": f"{max_model} 约 ¥{max_cost:.4f}",
            "recommended_cost": recommended_cost,
            "max_cost": max_cost,
            "saved_cost": 0.0,
            "saving": 0.0,
        }
    if max_cost > 0 and recommended_cost > 0:
        saving = max(0.0, 1 - recommended_cost / max_cost)
        saved_cost = max(0.0, max_cost - recommended_cost)
        return {
            "summary": f"省 {saving:.0%}",
            "detail": f"{max_model} 约 ¥{max_cost:.4f}，少花 ¥{saved_cost:.4f}",
            "recommended_cost": recommended_cost,
            "max_cost": max_cost,
            "saved_cost": saved_cost,
            "saving": saving,
        }
    stats = MODEL_BENCHMARKS.get(model_name)
    most_expensive = MODEL_BENCHMARKS[max_model]
    if stats:
        saving = max(0.0, 1 - stats.cost_cny / most_expensive.cost_cny)
        return {
            "summary": f"省 {saving:.0%}",
            "detail": f"参考全量评测：{max_model} 成本约 ¥{most_expensive.cost_cny:.2f}",
            "recommended_cost": recommended_cost,
            "max_cost": max_cost,
            "saved_cost": 0.0,
            "saving": saving,
        }
    return {
        "summary": "待确认",
        "detail": "缺少可用价格信息",
        "recommended_cost": recommended_cost,
        "max_cost": max_cost,
        "saved_cost": 0.0,
        "saving": 0.0,
    }


def _cost_saving_vs_most_expensive(model_name: str, decision: RoutingDecision) -> float | None:
    stats = MODEL_BENCHMARKS.get(model_name)
    most_expensive = MODEL_BENCHMARKS["qwen3.7-max"]
    if stats:
        if model_name == "qwen3.7-max":
            return 0.0
        return max(0.0, 1 - stats.cost_cny / most_expensive.cost_cny)
    if decision.baseline_cost > 0 and decision.estimated_cost > 0:
        return max(0.0, 1 - decision.estimated_cost / decision.baseline_cost)
    return None


def _extract_json_object(text: str) -> dict[str, Any] | None:
    cleaned = text.strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.strip("`")
        if cleaned.lower().startswith("json"):
            cleaned = cleaned[4:].strip()
    try:
        parsed = json.loads(cleaned)
        return parsed if isinstance(parsed, dict) else None
    except json.JSONDecodeError:
        pass

    decoder = json.JSONDecoder()
    for match in re.finditer(r"\{", cleaned):
        try:
            parsed, _ = decoder.raw_decode(cleaned[match.start() :])
            return parsed if isinstance(parsed, dict) else None
        except json.JSONDecodeError:
            continue
    return None


def _strip_thinking_text(text: str) -> str:
    without_blocks = re.sub(r"<think>.*?</think>", "", text, flags=re.DOTALL | re.IGNORECASE)
    without_open_block = re.sub(r"<think>.*", "", without_blocks, flags=re.DOTALL | re.IGNORECASE)
    return without_open_block.strip()


def _clean_review_text(text: str) -> str:
    cleaned = _strip_thinking_text(text)
    cleaned = re.sub(r"```(?:json)?", "", cleaned, flags=re.IGNORECASE).replace("```", "")
    return cleaned.strip()


def _review_model_name() -> str:
    registry = _model_registry()
    if registry.get_by_name(RECOMMENDATION_REVIEW_MODEL):
        return RECOMMENDATION_REVIEW_MODEL
    return RECOMMENDATION_REVIEW_FALLBACK_MODEL


def _review_recommendation(
    *,
    task: str,
    decision: RoutingDecision,
    recommended_model: str,
    treatment: str,
    api_key: str,
) -> dict[str, Any]:
    judge_model = _review_model_name()
    model = _model_entry(judge_model)
    allowed_models = [row["模型"] for row in _catalog_rows()]
    allowed_treatments = [
        "快速处理",
        "稳妥处理",
        "多模态处理",
        "先计划再执行",
        "先计划、执行后复核",
        "需要确认后执行",
    ]
    prompt = {
        "task": task,
        "local_router_recommendation": {
            "model": recommended_model,
            "treatment": treatment,
            "category": _friendly_category(decision.task_features.category),
            "complexity": _friendly_complexity(decision.task_features.complexity),
            "verifiability": _friendly_verifiability(decision.task_features.verifiability),
            "requires_human_review": decision.requires_human_review,
            "strategy": decision.recommended_strategy.value,
            "analysis_mode": _analysis_mode_label(decision),
        },
        "allowed_model_names": allowed_models,
        "allowed_treatments": allowed_treatments,
        "questions_to_check": [
            "这个任务真的简单还是复杂？",
            "是否漏了风险？",
            "推荐模型是否过低或过高？",
            "是否需要先澄清？",
            "是否建议人工确认？",
        ],
    }
    messages = [
        {
            "role": "system",
            "content": (
                "你是 AirRouter 的推荐复核员。请只复核路由推荐，不要执行用户任务。"
                "你要判断本地 Router 是否低估难度、漏掉风险、过度使用昂贵模型，"
                "并给出可给普通用户看的中文解释。严格输出 JSON。"
                "不要输出 <think>，不要输出推理过程，不要输出 Markdown。"
            ),
        },
        {
            "role": "user",
            "content": (
                "请复核以下推荐方案，只输出一个 JSON 对象。"
                "suggested_model 必须为空字符串或来自 allowed_model_names，不能发明新模型名。\n"
                "suggested_treatment 必须为空字符串或来自 allowed_treatments，不能输出长句。\n"
                "{\n"
                '  "verdict": "agree|adjust|clarify|human_review",\n'
                '  "confidence": 0.0,\n'
                '  "suggested_model": "模型名或空字符串",\n'
                '  "suggested_treatment": "处理方式或空字符串",\n'
                '  "needs_clarification": false,\n'
                '  "needs_human_review": false,\n'
                '  "risk_notes": ["普通用户看得懂的注意点"],\n'
                '  "reason": "一句话说明"\n'
                "}\n\n"
                f"{json.dumps(prompt, ensure_ascii=False)}"
            ),
        },
    ]
    adapter = OpenAICompatAdapter(api_base=model.api_base, api_key=api_key)
    start = perf_counter()
    response = adapter.chat(
        judge_model,
        messages,
        temperature=0.0,
        max_tokens=4096,
    )
    elapsed_ms = int((perf_counter() - start) * 1000)
    parsed = _extract_json_object(response.content) or {}
    cost = _estimate_stream_cost(judge_model, response.input_tokens, response.output_tokens)
    reason = _clean_review_text(str(parsed.get("reason") or ""))
    if not parsed:
        reason = "本次复核没有返回可用的结构化建议，当前执行方案不会改变。你可以重试深度复核，或继续使用当前方案。"
    return {
        "model": judge_model,
        "raw": response.content,
        "verdict": str(parsed.get("verdict") or "parse_failed"),
        "confidence": parsed.get("confidence"),
        "suggested_model": str(parsed.get("suggested_model") or ""),
        "suggested_treatment": str(parsed.get("suggested_treatment") or ""),
        "needs_clarification": bool(parsed.get("needs_clarification", False)),
        "needs_human_review": bool(parsed.get("needs_human_review", False)),
        "risk_notes": [
            _clean_review_text(str(note))
            for note in parsed.get("risk_notes", [])
            if str(note).strip()
        ]
        if isinstance(parsed.get("risk_notes"), list)
        else [],
        "reason": reason,
        "input_tokens": response.input_tokens,
        "output_tokens": response.output_tokens,
        "latency_ms": response.latency_ms or elapsed_ms,
        "cost_cny": round(cost, 6),
    }


def _analysis_mode_label(decision: RoutingDecision) -> str:
    diag = decision.complexity_analysis
    if diag and diag.ml_available and diag.mode == "opensquilla":
        return "OpenSquilla ML final"
    if diag and diag.mode != "rules" and diag.fallback_reason:
        return "快速规则分析（ML 不可用）"
    return "快速规则分析"


def _dynamic_notices(task: str, decision: RoutingDecision) -> list[str]:
    text = task.lower()
    notices: list[str] = []
    category = decision.task_features.category

    finance_terms = ("金额", "付款", "报销", "发票", "预算", "账单", "对账", "财务", "合同", "支付")
    code_terms = ("代码", "bug", "sql", "数据库", "接口", "配置", "日志", "生产", "权限", "部署")
    personal_terms = ("身份证", "手机号", "薪资", "员工", "客户", "个人信息", "隐私")
    irreversible_terms = ("删除", "drop", "delete", "update", "权限", "生产", "发布", "付款", "转账")

    if any(term in text for term in finance_terms):
        notices.append("涉及金额、合同或正式结论时，建议先确认口径再使用结果。")
    if category in {
        TaskCategory.CODE_GENERATION,
        TaskCategory.CODE_REVIEW,
        TaskCategory.REFACTORING,
        TaskCategory.BUG_ANALYSIS,
        TaskCategory.TEST_GENERATION,
        TaskCategory.DEVOPS,
    } or any(term in text for term in code_terms):
        notices.append("涉及代码、配置或日志时，建议先看方案，再决定是否执行修改。")
    if any(term in text for term in personal_terms):
        notices.append("可能包含个人或客户信息，建议避免粘贴不必要的敏感内容。")
    if any(term in text for term in irreversible_terms) or decision.risk_level == RiskLevel.HIGH:
        notices.append("如果会影响正式业务、权限、数据或对外承诺，请确认后再执行。")
    if decision.requires_human_review:
        notices.append(decision.human_review_reason or "系统建议由你确认后再继续。")

    deduped = []
    for notice in notices:
        if notice not in deduped:
            deduped.append(notice)
    return deduped or ["暂未发现特别需要确认的事项，可以先按推荐方案执行。"]


def _why_this_model(model_name: str, treatment: str, decision: RoutingDecision) -> list[str]:
    stats = MODEL_BENCHMARKS.get(model_name)
    saving = _cost_saving_vs_most_expensive(model_name, decision)
    reasons = [
        f"系统判断这个任务适合「{treatment}」。",
        "推荐会优先在完成质量、成本和是否需要确认之间做平衡。",
    ]
    if stats:
        reasons.append(
            f"{model_name} 在最新路由评测中的通过率约 {stats.pass_rate:.1f}%，定位是「{stats.position}」。"
        )
    if saving and saving > 0.05:
        reasons.append(f"相比使用 qwen3.7-max，预计节省约 {saving:.0%} 成本。")
    return reasons


def _catalog_rows() -> list[dict[str, Any]]:
    rows = []
    for row in model_catalog():
        name = str(row.get("模型", ""))
        stats = MODEL_BENCHMARKS.get(name)
        rows.append(
            {
                "模型": name,
                "适合用在": MODEL_DESCRIPTIONS.get(name, "通用模型"),
                "评测通过率": f"{stats.pass_rate:.1f}%" if stats else "-",
                "相对定位": stats.position if stats else "-",
                "输入价（元/1K）": row.get("输入价（元/1K）"),
                "输出价（元/1K）": row.get("输出价（元/1K）"),
            }
        )
    return rows


def _render_message(role: str, content: str) -> None:
    if role == "user":
        st.markdown(
            f'<div class="ar-chat-message user">{_html_text(content)}</div>',
            unsafe_allow_html=True,
        )
        return

    st.markdown('<div class="ar-assistant-message">', unsafe_allow_html=True)
    st.markdown(content)
    st.markdown("</div>", unsafe_allow_html=True)


def _render_deep_review(
    review: dict[str, Any],
    *,
    initial_model: str,
    initial_treatment: str,
    current_model: str,
    current_treatment: str,
) -> None:
    verdict_map = {
        "agree": "复核同意当前推荐",
        "adjust": "复核建议调整方案",
        "clarify": "复核建议先补充信息",
        "human_review": "复核建议人工确认",
        "parse_failed": "复核未形成可用建议",
        "unclear": "复核结果需要人工判断",
    }
    verdict_key = str(review.get("verdict") or "")
    if verdict_key == "unclear" and "<think" in str(review.get("reason") or "").lower():
        verdict_key = "parse_failed"
    verdict = verdict_map.get(verdict_key, str(review.get("verdict") or "复核完成"))
    visible_reason = _clean_review_text(str(review.get("reason") or ""))
    if not visible_reason:
        visible_reason = "本次复核没有返回可用的结构化建议，当前执行方案不会改变。你可以重试深度复核，或继续使用当前方案。"
    st.markdown(
        f"""
        <div class="ar-review-card">
          <div class="ar-review-topline">深度复核 · {review.get("model", "")}</div>
          <div class="ar-review-title">{_html_text(verdict)}</div>
          <div class="ar-review-copy">{_html_text(visible_reason)}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    notes = [
        _clean_review_text(str(note))
        for note in (review.get("risk_notes") or [])
        if _clean_review_text(str(note))
    ]
    if notes:
        st.markdown("复核注意点：")
        for note in notes:
            st.markdown(f"- {note}")
    suggested_model = str(review.get("suggested_model") or "").strip()
    suggested_treatment = str(review.get("suggested_treatment") or "").strip()
    valid_models = {row["模型"] for row in _catalog_rows()}
    valid_treatments = {
        "快速处理",
        "稳妥处理",
        "多模态处理",
        "先计划再执行",
        "先计划、执行后复核",
        "需要确认后执行",
    }
    if suggested_model or suggested_treatment:
        shown_model = suggested_model if suggested_model in valid_models else current_model
        shown_treatment = suggested_treatment if suggested_treatment in valid_treatments else current_treatment
        st.caption(
            f"复核建议：{shown_model} · {shown_treatment}"
        )
        a1, a2, a3 = st.columns([1.1, 1.2, 1.1])
        if a1.button("采用复核建议", use_container_width=True):
            if suggested_model in valid_models:
                st.session_state["ar_selected_model"] = suggested_model
                st.session_state["ar_execution_model_picker"] = suggested_model
            if suggested_treatment in valid_treatments:
                st.session_state["ar_selected_treatment"] = suggested_treatment
            _save_current_conversation()
            st.rerun()
        if a2.button("不采用，保留当前方案", use_container_width=True):
            st.session_state["ar_judge_review"] = None
            _save_current_conversation()
            st.rerun()
        if a3.button("恢复初始推荐", use_container_width=True):
            st.session_state["ar_selected_model"] = initial_model
            st.session_state["ar_execution_model_picker"] = initial_model
            st.session_state["ar_selected_treatment"] = initial_treatment
            st.session_state["ar_judge_review"] = None
            _save_current_conversation()
            st.rerun()
    else:
        b1, b2 = st.columns([1, 1])
        if b1.button("知道了，保留当前方案", use_container_width=True):
            st.session_state["ar_judge_review"] = None
            _save_current_conversation()
            st.rerun()
        if b2.button("恢复初始推荐", use_container_width=True):
            st.session_state["ar_selected_model"] = initial_model
            st.session_state["ar_execution_model_picker"] = initial_model
            st.session_state["ar_selected_treatment"] = initial_treatment
            st.session_state["ar_judge_review"] = None
            _save_current_conversation()
            st.rerun()


def _render_recommendation(task: str, decision: RoutingDecision) -> None:
    attachments = list(st.session_state.get("ar_attachments", []))
    has_image = _has_image_attachments(attachments)
    initial_model = IMAGE_MODEL if has_image else _primary_model(decision)
    initial_treatment, initial_copy = _friendly_strategy(decision)
    if has_image:
        st.session_state["ar_selected_model"] = IMAGE_MODEL
        st.session_state["ar_manual_open"] = False
    recommended_model = IMAGE_MODEL if has_image else st.session_state.get("ar_selected_model") or initial_model
    treatment = st.session_state.get("ar_selected_treatment") or initial_treatment
    treatment_title, treatment_copy = initial_treatment, initial_copy
    if st.session_state.get("ar_selected_treatment"):
        treatment_title = treatment
        treatment_copy = "你已手动调整处理方式，系统会按当前选择继续。"
    if has_image and not st.session_state.get("ar_selected_treatment"):
        treatment_title = "多模态处理"
        treatment_copy = "检测到图片附件，将使用支持图片的模型完成。"

    completion = _completion_label(recommended_model, decision)
    cost_compare = _cost_comparison_payload(recommended_model, decision)
    display_cost = _estimate_decision_cost_for_model(recommended_model, decision)
    notices = _dynamic_notices(task, decision)
    reasons = _why_this_model(recommended_model, treatment_title, decision)
    tool_items = _tool_status_items(
        attachments=attachments,
        report_tool=_should_use_report_tool("自动", task),
    )
    if has_image:
        notices.insert(0, "检测到图片附件，本次执行模型固定为 qwen3.7-plus；其他模型可能无法读取图片。")
        reasons.insert(0, "这次任务包含图片，系统会先走多模态路径，而不是普通文本路由。")
    mode_label = _analysis_mode_label(decision)

    st.markdown('<div class="ar-plan-card">', unsafe_allow_html=True)
    st.markdown(
        f"""
        <div class="ar-plan-topline">
          <span>推荐方案</span>
          <span>{mode_label}</span>
        </div>
        <div class="ar-plan-title">{recommended_model}</div>
        <div class="ar-plan-subtitle">{treatment_title} · {treatment_copy}</div>
        <div class="ar-plan-metrics">
          <div><span>完成把握</span><strong>{completion}</strong></div>
          <div><span>预计成本</span><strong>¥{display_cost:.4f}</strong></div>
          <div>
            <span>成本对比</span>
            <strong>{cost_compare["summary"]}</strong>
            <small>{cost_compare["detail"]}</small>
          </div>
        </div>
        <div class="ar-active-plan">
          <span>点击「接受并执行」将使用</span>
          <strong>{recommended_model}</strong>
          <em>{treatment_title}</em>
        </div>
        """,
        unsafe_allow_html=True,
    )
    _render_tool_strip(tool_items)

    st.markdown('<div class="ar-reason-grid">', unsafe_allow_html=True)
    st.markdown('<div><h4>为什么这么推荐</h4>', unsafe_allow_html=True)
    for reason in reasons:
        st.markdown(f"- {reason}")
    st.markdown("</div>", unsafe_allow_html=True)
    st.markdown('<div><h4>需要注意</h4>', unsafe_allow_html=True)
    for notice in notices:
        st.markdown(f"- {notice}")
    st.markdown("</div></div>", unsafe_allow_html=True)

    with st.expander("查看系统依据"):
        evidence_rows = [
            {"项目": "分析方式", "判断": mode_label},
            {"项目": "任务类型", "判断": _friendly_category(decision.task_features.category)},
            {"项目": "任务难度", "判断": _friendly_complexity(decision.task_features.complexity)},
            {"项目": "结果是否容易检查", "判断": _friendly_verifiability(decision.task_features.verifiability)},
            {"项目": "当前处理方式", "判断": treatment_title},
            {"项目": "内部策略", "判断": decision.recommended_strategy.value},
        ]
        diag = decision.complexity_analysis
        if diag and diag.ml_available and diag.ml_final_class:
            evidence_rows.append({"项目": "OpenSquilla 判断", "判断": diag.ml_final_class})
        elif diag and diag.fallback_reason:
            evidence_rows.append({"项目": "ML 回退原因", "判断": diag.fallback_reason})
        st.dataframe(evidence_rows, use_container_width=True, hide_index=True)

    review = st.session_state.get("ar_judge_review")
    if isinstance(review, dict):
        _render_deep_review(
            review,
            initial_model=initial_model,
            initial_treatment=initial_treatment,
            current_model=recommended_model,
            current_treatment=treatment_title,
        )

    c1, c2, c3, c4, c5 = st.columns([1.25, 1.05, .85, .85, 1])
    if c1.button("接受并执行", type="primary", use_container_width=True):
        api_key = str(st.session_state.get("api_key", ""))
        if not api_key:
            st.error("当前会话没有 API Key，请退出后重新进入。")
            st.stop()
        try:
            placeholder = st.empty()
            result = _execute_with_user_key(
                model_name=recommended_model,
                api_key=api_key,
                user_content=task,
                attachments=_image_attachments_only(attachments),
                placeholder=placeholder,
            )
        except AdapterError as exc:
            st.error(f"模型调用失败：{exc}")
            st.stop()
        except Exception as exc:
            st.error(f"执行失败：{exc}")
            st.stop()

        st.session_state["ar_selected_model"] = recommended_model
        st.session_state["ar_execution_model_picker"] = recommended_model
        st.session_state["ar_selected_treatment"] = treatment_title
        st.session_state["ar_execution_started"] = True
        st.session_state["ar_messages"].append({"role": "assistant", "content": result["content"]})
        st.session_state["ar_execution_records"].append(result)
        _save_current_conversation()
        st.rerun()
    if c2.button("深度复核", use_container_width=True):
        api_key = str(st.session_state.get("api_key", ""))
        if not api_key:
            st.error("当前会话没有 API Key，请退出后重新进入。")
            st.stop()
        try:
            with st.spinner("正在复核推荐方案..."):
                review_result = _review_recommendation(
                    task=task,
                    decision=decision,
                    recommended_model=recommended_model,
                    treatment=treatment_title,
                    api_key=api_key,
                )
        except AdapterError as exc:
            st.error(f"深度复核失败：{exc}")
            st.stop()
        except Exception as exc:
            st.error(f"深度复核失败：{exc}")
            st.stop()
        st.session_state["ar_judge_review"] = review_result
        _save_current_conversation()
        st.rerun()
    if c3.button("更省钱", use_container_width=True, disabled=has_image):
        st.session_state["ar_selected_model"] = "deepseek-v4-flash"
        st.session_state["ar_selected_treatment"] = "快速处理"
        st.session_state["ar_judge_review"] = None
        _save_current_conversation()
        st.rerun()
    if c4.button("更稳妥", use_container_width=True, disabled=has_image):
        st.session_state["ar_selected_model"] = "qwen3.7-plus"
        st.session_state["ar_selected_treatment"] = "稳妥处理"
        st.session_state["ar_judge_review"] = None
        _save_current_conversation()
        st.rerun()
    if c5.button("手动选择", use_container_width=True, disabled=has_image):
        st.session_state["ar_manual_open"] = not st.session_state.get("ar_manual_open", False)
        _save_current_conversation()
        st.rerun()

    st.markdown(
        '<div class="ar-review-cost-note">深度复核会调用大模型，点击后会消耗 token 并产生少量费用。</div>',
        unsafe_allow_html=True,
    )

    if st.session_state.get("ar_manual_open"):
        st.markdown("#### 手动调整")
        models = [row["模型"] for row in _catalog_rows()]
        selected = st.selectbox(
            "选择模型",
            models,
            index=models.index(recommended_model) if recommended_model in models else 0,
        )
        treatment_options = ["快速处理", "稳妥处理", "先计划再执行", "先计划、执行后复核", "需要确认后执行"]
        selected_treatment = st.selectbox(
            "选择处理方式",
            treatment_options,
            index=treatment_options.index(treatment_title)
            if treatment_title in treatment_options
            else 1,
        )
        st.dataframe(_catalog_rows(), use_container_width=True, hide_index=True)
        if st.button("应用调整", use_container_width=True):
            st.session_state["ar_selected_model"] = selected
            st.session_state["ar_selected_treatment"] = selected_treatment
            st.session_state["ar_manual_open"] = False
            st.session_state["ar_judge_review"] = None
            _save_current_conversation()
            st.rerun()
    st.markdown("</div>", unsafe_allow_html=True)


def _friendly_category(category: TaskCategory) -> str:
    return {
        TaskCategory.CODE_GENERATION: "代码生成",
        TaskCategory.CODE_REVIEW: "代码审查",
        TaskCategory.REFACTORING: "代码整理",
        TaskCategory.BUG_ANALYSIS: "问题分析",
        TaskCategory.DOCUMENTATION: "文档处理",
        TaskCategory.ARCHITECTURE: "方案设计",
        TaskCategory.TEST_GENERATION: "测试生成",
        TaskCategory.DATA_PROCESSING: "数据处理",
        TaskCategory.DEVOPS: "运维配置",
        TaskCategory.CONVERSATION: "通用问答",
    }.get(category, category.value)


def _friendly_complexity(complexity: ComplexityLevel) -> str:
    return {
        ComplexityLevel.SIMPLE: "简单",
        ComplexityLevel.MODERATE: "中等",
        ComplexityLevel.COMPLEX: "较复杂",
        ComplexityLevel.EXPERT: "专家级",
    }.get(complexity, complexity.value)


def _friendly_verifiability(level: VerifiabilityLevel) -> str:
    return {
        VerifiabilityLevel.HIGH: "容易检查",
        VerifiabilityLevel.MEDIUM: "部分可检查",
        VerifiabilityLevel.LOW: "主要依赖人工判断",
    }.get(level, level.value)


def _run_analysis(task: str, context: str) -> RoutingDecision:
    full_task = "\n".join(part for part in [task.strip(), context.strip()] if part).strip()
    complexity_context = ComplexityContext.from_task(
        full_task,
        history=tuple(),
    )
    config = load_all()
    bundle_dir = config.complexity.bundle_dir
    if bundle_dir is None and OPENSQUILLA_BUNDLE_PATH.exists():
        bundle_dir = OPENSQUILLA_BUNDLE_PATH
    ml_config = config.complexity.model_copy(
        update={
            "mode": "opensquilla",
            "bundle_dir": bundle_dir,
            "require_runtime": False,
        }
    )
    config = config.model_copy(update={"complexity": ml_config})
    return Router(config=config).route(full_task, complexity_context=complexity_context)


def _model_registry() -> ModelRegistry:
    return ModelRegistry(load_all().models)


def _model_entry(model_name: str) -> Any:
    model = _model_registry().get_by_name(model_name)
    if model is None:
        raise ValueError(f"未知模型：{model_name}")
    return model


def _estimate_tokens(text: str) -> int:
    return max(1, len(text) // 4)


def _estimate_stream_cost(model_name: str, input_tokens: int, output_tokens: int) -> float:
    model = _model_entry(model_name)
    return (
        input_tokens / 1000 * model.input_price_per_1k
        + output_tokens / 1000 * model.output_price_per_1k
    )


def _estimate_decision_cost_for_model(model_name: str, decision: RoutingDecision) -> float:
    try:
        return _estimate_stream_cost(
            model_name,
            decision.task_features.estimated_input_tokens,
            decision.task_features.estimated_output_tokens,
        )
    except Exception:
        return decision.estimated_cost


def _history_as_text(messages: list[dict[str, str]]) -> str:
    parts = []
    for message in messages:
        role = "用户" if message.get("role") == "user" else "助手"
        parts.append(f"{role}：{message.get('content', '')}")
    return "\n\n".join(parts)


def _compress_context_if_needed(api_key: str) -> None:
    history = list(st.session_state.get("ar_execution_messages", []))
    if len(history) <= CONTEXT_KEEP_RECENT_MESSAGES:
        return

    existing_summary = str(st.session_state.get("ar_context_summary", ""))
    total_chars = len(existing_summary) + sum(len(str(m.get("content", ""))) for m in history)
    if total_chars <= CONTEXT_COMPRESSION_CHAR_LIMIT:
        return

    older = history[:-CONTEXT_KEEP_RECENT_MESSAGES]
    recent = history[-CONTEXT_KEEP_RECENT_MESSAGES:]
    model = _model_entry(CONTEXT_SUMMARY_MODEL)
    prompt = (
        "请把下面这段企业助手对话压缩成后续模型继续执行任务所需的上下文记忆。\n"
        "要求：保留用户目标、已做决策、关键事实、约束、未完成事项、用户偏好；"
        "删除寒暄和重复内容；不要新增事实；用中文分点输出，控制在 1200 字以内。\n\n"
        f"已有摘要：\n{existing_summary or '无'}\n\n"
        f"需要压缩的较早对话：\n{_history_as_text(older)}"
    )
    try:
        adapter = OpenAICompatAdapter(api_base=model.api_base, api_key=api_key)
        response = adapter.chat(
            CONTEXT_SUMMARY_MODEL,
            [
                {"role": "system", "content": "你负责压缩对话上下文，只保留可执行记忆。"},
                {"role": "user", "content": prompt},
            ],
            temperature=0.0,
            max_tokens=1600,
        )
    except (AdapterAuthError, AdapterRateLimitError, AdapterTimeoutError, AdapterError) as exc:
        st.warning(f"上下文自动整理失败，本次会保留原始上下文继续：{exc}")
        return

    st.session_state["ar_context_summary"] = response.content.strip()
    st.session_state["ar_context_summary_updated_at"] = datetime.now().strftime("%H:%M")
    st.session_state["ar_execution_messages"] = recent
    st.session_state.setdefault("ar_execution_records", []).append(
        {
            "model": CONTEXT_SUMMARY_MODEL,
            "content": "上下文已自动整理。",
            "input_tokens": response.input_tokens,
            "output_tokens": response.output_tokens,
            "latency_ms": response.latency_ms or 0,
            "cost_cny": round(
                _estimate_stream_cost(
                    CONTEXT_SUMMARY_MODEL,
                    response.input_tokens,
                    response.output_tokens,
                ),
                6,
            ),
            "kind": "context_compression",
        }
    )


def _build_execution_messages(
    user_content: str,
    attachments: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    attachment_items = attachments or []
    messages: list[dict[str, Any]] = [{"role": "system", "content": SYSTEM_PROMPT}]
    summary = str(st.session_state.get("ar_context_summary", "")).strip()
    if summary:
        messages.append(
            {
                "role": "system",
                "content": f"以下是较早对话的压缩记忆，请在后续回答中保持一致：\n{summary}",
            }
        )
    messages.extend(
        {
            "role": str(message.get("role", "user")),
            "content": str(message.get("content", "")),
        }
        for message in st.session_state.get("ar_execution_messages", [])
    )
    messages.append(
        {
            "role": "user",
            "content": _build_multimodal_user_content(user_content, attachment_items),
        }
    )
    return messages


def _openai_error_message(exc: Exception) -> str:
    if isinstance(exc, AuthenticationError):
        return "API Key 无效、过期或没有调用该模型的权限。"
    if isinstance(exc, RateLimitError):
        return "触发了接口限流，请稍后重试或降低并发。"
    if isinstance(exc, APITimeoutError):
        return "模型响应超时，请稍后重试。"
    if isinstance(exc, APIError):
        return f"模型接口返回错误：{exc}"
    return str(exc)


def _stream_model_response(
    *,
    model_name: str,
    api_key: str,
    messages: list[dict[str, Any]],
    placeholder: Any,
) -> dict[str, Any]:
    model = _model_entry(model_name)
    client = OpenAI(base_url=model.api_base, api_key=api_key)
    request_kwargs: dict[str, Any] = {
        "model": model_name,
        "messages": messages,
        "max_tokens": 4096,
        "temperature": 0.2,
        "stream": True,
    }
    if model_name.lower().startswith("kimi-"):
        request_kwargs.pop("temperature", None)

    content = ""
    start = perf_counter()
    try:
        stream = client.chat.completions.create(**request_kwargs)
        for chunk in stream:
            if not chunk.choices:
                continue
            delta = chunk.choices[0].delta.content or ""
            if delta:
                content += delta
            elapsed = int(perf_counter() - start)
            with placeholder.container():
                st.markdown('<div class="ar-stream-box">', unsafe_allow_html=True)
                st.markdown(
                    f'<div class="ar-thinking">已思考 {elapsed}s</div>',
                    unsafe_allow_html=True,
                )
                st.markdown(content or "正在生成回答...")
                st.markdown("</div>", unsafe_allow_html=True)
    except (AuthenticationError, RateLimitError, APITimeoutError, APIError) as exc:
        raise AdapterError(_openai_error_message(exc)) from exc

    elapsed_ms = int((perf_counter() - start) * 1000)
    input_tokens = _estimate_tokens("\n".join(_message_text_for_estimation(m) for m in messages))
    output_tokens = _estimate_tokens(content)
    cost = _estimate_stream_cost(model_name, input_tokens, output_tokens)
    return {
        "model": model_name,
        "content": content.strip() or "模型没有返回内容。",
        "input_tokens": input_tokens,
        "output_tokens": output_tokens,
        "latency_ms": elapsed_ms,
        "cost_cny": round(cost, 6),
        "kind": "model_execution",
    }


def _execute_with_user_key(
    *,
    model_name: str,
    api_key: str,
    user_content: str,
    attachments: list[dict[str, Any]] | None = None,
    placeholder: Any,
) -> dict[str, Any]:
    attachment_items = attachments or []
    if _has_image_attachments(attachment_items):
        model_name = IMAGE_MODEL
    _compress_context_if_needed(api_key)
    messages = _build_execution_messages(user_content, attachment_items)
    result = _stream_model_response(
        model_name=model_name,
        api_key=api_key,
        messages=messages,
        placeholder=placeholder,
    )
    user_history_content = _build_task_with_materials(user_content, "", attachment_items)
    st.session_state.setdefault("ar_execution_messages", []).extend(
        [
            {"role": "user", "content": user_history_content},
            {"role": "assistant", "content": result["content"]},
        ]
    )
    return result


configure_page("智能工作台")
user_context = require_login()
_init_state()
_render_conversation_sidebar()

st.markdown('<div class="ar-workspace">', unsafe_allow_html=True)

st.markdown('<div class="ar-chat-header">AirRouter</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="ar-chat-subtitle">描述任务，先看推荐方案，确认后再执行。</div>',
    unsafe_allow_html=True,
)
if st.session_state.get("ar_context_summary"):
    updated_at = st.session_state.get("ar_context_summary_updated_at") or "刚刚"
    st.caption(f"上下文较长，已自动整理 · {updated_at}")

st.markdown('<div class="ar-chat-thread">', unsafe_allow_html=True)
for message in st.session_state["ar_messages"]:
    _render_message(str(message["role"]), str(message["content"]))
st.markdown("</div>", unsafe_allow_html=True)

current_analysis = st.session_state.get("ar_analysis")
if current_analysis and not st.session_state.get("ar_execution_started"):
    _render_recommendation(current_analysis["task"], current_analysis["decision"])

if st.session_state.get("ar_execution_started"):
    model_names = [row["模型"] for row in _catalog_rows()]
    active_attachments = list(st.session_state.get("ar_attachments", []))
    has_active_image = _has_image_attachments(active_attachments)
    saved_model = IMAGE_MODEL if has_active_image else st.session_state.get("ar_selected_model") or "qwen3.7-plus"
    if saved_model not in model_names:
        saved_model = "qwen3.7-plus"
    if st.session_state.get("ar_execution_model_picker") not in model_names:
        st.session_state["ar_execution_model_picker"] = saved_model
    if has_active_image:
        st.session_state["ar_execution_model_picker"] = IMAGE_MODEL
    model = str(st.session_state.get("ar_execution_model_picker") or saved_model)
    st.session_state["ar_selected_model"] = model
    treatment = st.session_state.get("ar_selected_treatment") or "稳妥处理"
    st.markdown(
        f"""
        <div class="ar-execution-bar">
          <strong>{model}</strong>
          <span>{treatment} · 上下文会随模型切换保留</span>
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.markdown('<div class="ar-input-shell">', unsafe_allow_html=True)
    follow_up_upload_version = int(st.session_state.get("ar_follow_up_upload_version", 0))
    follow_up = st.text_area(
        "继续输入",
        placeholder="继续追问、补充要求，或点左下角 + 添加文件...",
        height=104,
        label_visibility="collapsed",
        key=f"ar_follow_up_text_{follow_up_upload_version}",
    )
    st.markdown(
        '<div class="ar-composer-hint">上下文会保留；切换模型只影响下一次回答。</div>',
        unsafe_allow_html=True,
    )
    if active_attachments:
        st.caption(f"已保留 {len(active_attachments)} 个资料，可继续引用；新增文件请点左下角 +。")
    file_col, model_col, send_col = st.columns([1.0, 1.25, 1.0])
    with file_col:
        follow_up_files, follow_up_context = _render_material_popover(
            f"ar_follow_up_{follow_up_upload_version}",
            retained_attachments=active_attachments,
        )
    with model_col:
        model = st.selectbox(
            "执行模型",
            model_names,
            key="ar_execution_model_picker",
            disabled=has_active_image,
            label_visibility="collapsed",
        )
    with send_col:
        submitted_follow_up = st.button("发送", type="primary", use_container_width=True)
    if has_active_image:
        st.caption("当前对话包含图片，执行模型固定为 qwen3.7-plus。")
    follow_up_report_tool = _should_use_report_tool("自动", follow_up, follow_up_context)
    if follow_up_files:
        preview_attachments, preview_warnings = _process_uploaded_files(list(follow_up_files or []))
        _render_attachment_chips(preview_attachments)
        for warning in preview_warnings:
            st.caption(f"附件提示：{warning}")
    _render_tool_strip(
        _tool_status_items(
            uploaded_count=len(follow_up_files or []),
            report_tool=follow_up_report_tool,
        )
    )
    st.markdown("</div>", unsafe_allow_html=True)
    st.session_state["ar_selected_model"] = model
    if submitted_follow_up:
        current = st.session_state.get("ar_analysis") or {}
        original_task = str(current.get("task", ""))
        api_key = str(st.session_state.get("api_key", ""))
        if not api_key:
            st.error("当前会话没有 API Key，请退出后重新登录。")
            st.stop()
        user_content = "\n\n".join(
            part for part in [follow_up.strip(), follow_up_context.strip()] if part
        )
        display_user_content = user_content
        new_attachments, attachment_warnings = _process_uploaded_files(list(follow_up_files or []))
        execution_attachments = list(st.session_state.get("ar_attachments", [])) + new_attachments
        report_tool = _should_use_report_tool("自动", user_content)
        if report_tool:
            user_content = "\n\n".join(
                part for part in [user_content, _report_tool_instruction()] if part
            )
        if len(execution_attachments) > MAX_ATTACHMENTS:
            attachment_warnings.append(f"当前对话最多保留 {MAX_ATTACHMENTS} 个附件，已保留最新附件。")
            execution_attachments = execution_attachments[-MAX_ATTACHMENTS:]
        st.session_state["ar_attachments"] = execution_attachments
        active_tools = _tool_status_items(
            attachments=new_attachments,
            report_tool=report_tool,
        )
        if active_tools:
            _render_tool_strip(active_tools, active=True)
        if _has_image_attachments(execution_attachments):
            model = IMAGE_MODEL
            st.session_state["ar_execution_model_picker"] = IMAGE_MODEL
            st.session_state["ar_selected_model"] = IMAGE_MODEL
        if not user_content:
            user_content = (
                "请根据新上传的附件继续处理。"
                if new_attachments
                else
                "请继续处理上面的任务。"
                if st.session_state.get("ar_execution_messages")
                else original_task
            )
        try:
            placeholder = st.empty()
            result = _execute_with_user_key(
                model_name=model,
                api_key=api_key,
                user_content=user_content,
                attachments=execution_attachments,
                placeholder=placeholder,
            )
        except AdapterError as exc:
            st.error(f"模型调用失败：{exc}")
            st.stop()
        except Exception as exc:
            st.error(f"执行失败：{exc}")
            st.stop()

        if follow_up.strip() or follow_up_context.strip() or new_attachments:
            st.session_state["ar_messages"].append(
                {
                    "role": "user",
                    "content": _build_user_display_text(display_user_content, new_attachments),
                }
            )
        tool_message = _tool_status_message(active_tools, active=True)
        if tool_message:
            st.session_state["ar_messages"].append(
                {
                    "role": "assistant",
                    "content": tool_message,
                }
            )
        if attachment_warnings:
            st.session_state["ar_messages"].append(
                {
                    "role": "assistant",
                    "content": "附件提示：\n" + "\n".join(f"- {item}" for item in attachment_warnings),
                }
            )
        st.session_state["ar_messages"].append({"role": "assistant", "content": result["content"]})
        st.session_state["ar_execution_records"].append(result)
        st.session_state["ar_follow_up_upload_version"] = follow_up_upload_version + 1
        _save_current_conversation()
        st.rerun()

if not current_analysis and not st.session_state.get("ar_execution_started"):
    st.markdown('<div class="ar-input-shell">', unsafe_allow_html=True)
    task_text = st.text_area(
        "输入你的问题",
        placeholder="把你想处理的事情发给我，比如审核报销、分析 SQL 风险、排查日志、整理文档...",
        height=132,
        label_visibility="collapsed",
        key="ar_initial_task_text",
    )
    st.markdown(
        '<div class="ar-composer-hint">可以直接提问，也可以点 + 添加 PDF、表格、图片或补充背景。</div>',
        unsafe_allow_html=True,
    )
    file_col, mode_col, submit_col = st.columns([1.05, 1.05, 1.0])
    with file_col:
        uploaded_files, extra_context = _render_material_popover("ar_initial")
    with mode_col:
        processing_mode = st.selectbox(
            "处理方式",
            PROCESSING_MODES,
            key="ar_processing_mode",
            label_visibility="collapsed",
        )
    with submit_col:
        submitted = st.button("分析", type="primary", use_container_width=True)
    preview_report_tool = _should_use_report_tool(
        processing_mode,
        task_text,
        extra_context,
    )
    if uploaded_files:
        preview_attachments, preview_warnings = _process_uploaded_files(list(uploaded_files or []))
        _render_attachment_chips(preview_attachments)
        for warning in preview_warnings:
            st.caption(f"附件提示：{warning}")
    _render_tool_strip(
        _tool_status_items(
            uploaded_count=len(uploaded_files or []),
            report_tool=preview_report_tool,
        )
    )
    st.markdown("</div>", unsafe_allow_html=True)
    if submitted:
        mode_instruction = _processing_mode_instruction(processing_mode)
        report_tool = _should_use_report_tool(
            processing_mode,
            task_text,
            extra_context,
        )
        if not task_text.strip() and not extra_context.strip() and not uploaded_files:
            st.error("先输入你想处理的问题，或上传/粘贴一份材料。")
            st.stop()
        if not task_text.strip():
            task_text = "请根据我提供的材料继续处理。"
        api_key = str(st.session_state.get("api_key", ""))
        attachments, attachment_warnings = _process_uploaded_files(list(uploaded_files or []))
        active_tools = _tool_status_items(attachments=attachments, report_tool=report_tool)
        if active_tools:
            _render_tool_strip(active_tools, active=True)
        combined_context = "\n\n".join(
            part
            for part in [
                extra_context.strip(),
                _report_tool_instruction() if report_tool else mode_instruction,
            ]
            if part
        )
        analysis_task = _build_task_with_materials(task_text, combined_context, attachments)
        try:
            decision_result = _run_analysis(analysis_task, "")
        except Exception as exc:
            st.error(f"分析失败：{exc}")
            st.stop()
        st.session_state["ar_selected_model"] = None
        st.session_state["ar_selected_treatment"] = None
        st.session_state["ar_execution_started"] = False
        st.session_state["ar_manual_open"] = False
        st.session_state["ar_execution_records"] = []
        st.session_state["ar_execution_messages"] = []
        st.session_state["ar_context_summary"] = ""
        st.session_state["ar_context_summary_updated_at"] = ""
        st.session_state["ar_judge_review"] = None
        st.session_state["ar_attachments"] = attachments
        st.session_state.pop("ar_execution_model_picker", None)
        st.session_state["ar_analysis"] = {
            "task": analysis_task,
            "decision": decision_result,
        }
        st.session_state["ar_messages"].append(
            {"role": "user", "content": _build_user_display_text(task_text, attachments)}
        )
        tool_message = _tool_status_message(active_tools, active=True)
        if tool_message:
            st.session_state["ar_messages"].append(
                {
                    "role": "assistant",
                    "content": tool_message,
                }
            )
        if attachment_warnings:
            st.session_state["ar_messages"].append(
                {
                    "role": "assistant",
                    "content": "附件提示：\n" + "\n".join(f"- {item}" for item in attachment_warnings),
                }
            )
        st.session_state["ar_messages"].append(
            {"role": "assistant", "content": "我先分析一下任务，再给你一个推荐方案。"}
        )
        _update_conversation_title(task_text, api_key)
        _save_current_conversation()
        st.rerun()

st.markdown("</div>", unsafe_allow_html=True)
